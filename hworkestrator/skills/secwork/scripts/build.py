#!/usr/bin/env python3
"""Build an information-security assignment (.docx and/or .pdf) from a solution JSON, with the fixed GTU header.

usage: build.py SOLUTION.json --week 3 --course "COURSE" [--out FILE.docx] [--pdf FILE.pdf] [--run]
                              [--student NAME --group GROUP]   (default: $STUDENT_NAME / $STUDENT_GROUP)

SOLUTION.json:
    {"title": "სახელისა და გვარის დაშიფვრა და ანალიზი",
     "tasks": [{"n": 1, "title": "სახელისა და გვარის დაშიფვრა",
                "condition": ["paragraph", ...],          # task statement + keys / method / formulas
                "code": "python source" | "",             # empty for a theory-only task
                "caption": "სახელისა და გვარის დაშიფვრის შედეგი",
                "answer": ["paragraph", ...]}],            # explanation / theory answer (optional)
     "conclusions": ["paragraph", ...]}
--run executes every task's code (python, 60 s limit) and stores the REAL stdout in "output" (written back
into the JSON, so a later build without --run keeps the same numbers even if the code is random).
Exit code 1 if any code failed; the failing task and its traceback are printed.
--pdf also typesets the same content with XeLaTeX (needs xelatex + the FreeSerif/FreeMono fonts).
"""
import argparse, datetime, json, os, shutil, subprocess, sys, tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "..", "assets")
BLUE = RGBColor(0x2F, 0x54, 0x96)
NAVY = RGBColor(0x1F, 0x38, 0x64)
TIMEOUT = 60


# ---------- running the code ----------
def run_code(code):
    """-> (ok, output). Runs in an empty temp dir; output = stdout (+ stderr on failure)."""
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "task.py")
        open(path, "w", encoding="utf8").write(code)
        try:
            r = subprocess.run([sys.executable, "-I", "-X", "utf8", path], cwd=d, capture_output=True, text=True,
                               timeout=TIMEOUT, stdin=subprocess.DEVNULL)
        except subprocess.TimeoutExpired:
            return False, f"timeout: the program ran longer than {TIMEOUT} s (no input() allowed)"
    out = r.stdout.rstrip()
    if r.returncode:
        return False, (out + "\n" + r.stderr.strip()).strip()
    return True, out


def run_all(sol):
    bad = 0
    for t in sol["tasks"]:
        if not t.get("code", "").strip():
            continue
        ok, out = run_code(t["code"])
        t["output"] = out
        print(f"=== task {t['n']} {'OK' if ok else 'FAILED'}\n{out[-3000:]}")
        bad += not ok
    return bad


def paras(v):
    """A paragraph list; tolerates a single string."""
    return [v] if isinstance(v, str) else [str(x) for x in v or []]


def task_title(t):
    return f"ამოცანა {t['n']} – {t.get('title', '')}".rstrip(" –")


# ---------- docx ----------
def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
    pr.append(shd)


def borders(table, color="BFBFBF"):
    pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"), e.set(qn("w:sz"), "6"), e.set(qn("w:space"), "0"), e.set(qn("w:color"), color)
        b.append(e)
    pr.append(b)


def keep_next(p):
    p.paragraph_format.keep_with_next = True
    return p


def label(doc, text):
    p = keep_next(doc.add_paragraph())
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(4)
    r = p.add_run(text)
    r.bold, r.font.color.rgb, r.font.size = True, NAVY, Pt(10.5)


def text(doc, s):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(s).font.size = Pt(11)
    return p


def mono_box(doc, s, fill):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    cell = t.rows[0].cells[0]
    shade(cell, fill)
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    for line in (s or "").rstrip("\n").split("\n"):
        p = cell.add_paragraph()
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line.replace("\t", "    "))
        r.font.name, r.font.size = "Consolas", Pt(8.5)
        r._element.rPr.rFonts.set(qn("w:cs"), "Sylfaen"), r._element.rPr.rFonts.set(qn("w:eastAsia"), "Sylfaen")
    if not cell.paragraphs:
        cell.add_paragraph()


def header(doc, student, group, course, number, date):
    doc.add_paragraph().add_run().add_picture(os.path.join(ASSETS, "logo_banner.png"), width=Cm(12))
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    left, right = t.rows[0].cells
    left.width, right.width = Cm(4.2), Cm(11.8)
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    left.paragraphs[0].add_run().add_picture(os.path.join(ASSETS, "logo_seal.png"), width=Cm(3.5))
    rows = [("სტუდენტი", student), ("ჯგუფი", group), ("კურსის დასახელება", course),
            ("დავალების ნომერი", number), ("თარიღი", date)]
    right.paragraphs[0]._p.getparent().remove(right.paragraphs[0]._p)
    for k, v in rows:
        p = right.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.4)
        r = p.add_run(f"{k}: ")
        r.bold, r.font.color.rgb, r.font.size = True, NAVY, Pt(10.5)
        p.add_run(v).font.size = Pt(10.5)


def heading(doc, s, first):
    p = keep_next(doc.add_paragraph())
    p.paragraph_format.space_before = Pt(56 if first else 0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.page_break_before = not first
    r = p.add_run(s)
    r.font.size, r.font.color.rgb = Pt(17), BLUE
    pPr = p._p.get_or_add_pPr()
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), "0")
    pPr.append(lvl)


def build_docx(sol, out, student, group, course, number, date, lang="Python"):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Sylfaen", Pt(11)
    normal.element.rPr.rFonts.set(qn("w:cs"), "Sylfaen")
    header(doc, student, group, course, number, date)
    fig = 0
    for i, t in enumerate(sol["tasks"]):
        heading(doc, task_title(t), i == 0)
        label(doc, "ამოცანის პირობა")
        for s in paras(t.get("condition")):
            text(doc, s)
        if t.get("code", "").strip():
            label(doc, f"პროგრამული კოდი (Source Code) – {lang}")
            mono_box(doc, t["code"], "F2F2F2")
            label(doc, "შედეგი / აუთფუთი (Output / Screenshot)")
            mono_box(doc, t.get("output", ""), "FFFFFF")
            fig += 1
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(f"სურ. {fig}. {t.get('caption') or t.get('title', '')}")
            r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), RGBColor(0x6B, 0x72, 0x80)
        if paras(t.get("answer")):
            label(doc, "პასუხი" if not t.get("code", "").strip() else "ახსნა")
            for s in paras(t["answer"]):
                text(doc, s)
    if paras(sol.get("conclusions")):
        label(doc, "დასკვნები")
        for s in paras(sol["conclusions"]):
            text(doc, s)
    doc.save(out)


# ---------- pdf (XeLaTeX, same layout as the docx) ----------
def tex_escape(t):  # single pass so replacements are not escaped again
    table = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
             "{": r"\{", "}": r"\}", "^": r"\^{}", "~": r"\textasciitilde{}"}
    return "".join(table.get(ch, ch) for ch in str(t))


def build_pdf(sol, out, student, group, course, number, date, lang="Python"):
    """Typeset with XeLaTeX into `out` (.pdf); the .tex is kept next to it."""
    out = os.path.abspath(out)
    work = tempfile.mkdtemp()
    os.makedirs(os.path.join(work, "latex_assets"))
    for f in ("logo_banner.png", "logo_seal.png"):
        shutil.copy(os.path.join(ASSETS, f), os.path.join(work, "latex_assets", f))
    defs = "".join("\\newcommand{\\%s}{%s}\n" % (k, tex_escape(v)) for k, v in (
        ("assignmentstudent", student), ("assignmentgroup", group), ("assignmentcourse", course),
        ("assignmentnumber", number), ("assignmentdate", date)))
    body, fig = [], 0
    for i, t in enumerate(sol["tasks"]):
        if i:
            body.append("\\clearpage")
        body.append("\\task{%s}" % tex_escape(task_title(t)))
        body.append("\\partlabel{ამოცანის პირობა}")
        body += [tex_escape(s) + "\n" for s in paras(t.get("condition"))]
        if t.get("code", "").strip():
            for kind, src in (("code", t["code"]), ("out", t.get("output", "") or " ")):
                open(os.path.join(work, f"{kind}{t['n']}.txt"), "w", encoding="utf8").write(
                    src.rstrip("\n").replace("\t", "    ") + "\n")
            body.append("\\partlabel{პროგრამული კოდი (Source Code) – %s}" % lang)
            body.append("\\codefile{code%s.txt}{codebg}" % t["n"])
            body.append("\\partlabel{შედეგი / აუთფუთი (Output / Screenshot)}")
            body.append("\\codefile{out%s.txt}{white}" % t["n"])
            fig += 1
            body.append("\\figcaption{%s}" % tex_escape(f"სურ. {fig}. {t.get('caption') or t.get('title', '')}"))
        if paras(t.get("answer")):
            body.append("\\partlabel{%s}" % ("ახსნა" if t.get("code", "").strip() else "პასუხი"))
            body += [tex_escape(s) + "\n" for s in paras(t["answer"])]
    if paras(sol.get("conclusions")):
        body.append("\\partlabel{დასკვნები}")
        body += [tex_escape(s) + "\n" for s in paras(sol["conclusions"])]
    pre = open(os.path.join(ASSETS, "preamble.tex"), encoding="utf8").read()
    tex = pre + "\n" + defs + "\\begin{document}\n\\assignmentheader\n\n" + "\n".join(body) + "\n\\end{document}\n"
    open(os.path.join(work, "doc.tex"), "w", encoding="utf8").write(tex)
    for _ in range(2):
        r = subprocess.run(["xelatex", "-interaction=nonstopmode", "-halt-on-error", "doc.tex"],
                           cwd=work, capture_output=True, text=True)
        if r.returncode:
            sys.exit("XeLaTeX FAILED (files kept in %s):\n%s" % (work, r.stdout[-2500:]))
    shutil.copy(os.path.join(work, "doc.pdf"), out)
    shutil.copy(os.path.join(work, "doc.tex"), os.path.splitext(out)[0] + ".tex")
    missing = sorted({l.split("There is no ")[1].split(" ")[0]
                      for l in open(os.path.join(work, "doc.log"), errors="ignore") if "Missing character" in l})
    if missing:
        print("warning: characters missing in the font:", " ".join(missing))
    shutil.rmtree(work, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("solution")
    ap.add_argument("--week", required=True, help="assignment number shown in the header")
    ap.add_argument("--out", help="the .docx (default: ~/Desktop/Assignment_<week>.docx)")
    ap.add_argument("--pdf", help="also typeset this .pdf with XeLaTeX")
    ap.add_argument("--run", action="store_true", help="execute the code and store the real outputs")
    ap.add_argument("--date", default=datetime.date.today().strftime("%d.%m.%Y"))
    ap.add_argument("--student", default=os.getenv("STUDENT_NAME", ""), help="default: $STUDENT_NAME")
    ap.add_argument("--group", default=os.getenv("STUDENT_GROUP", ""), help="default: $STUDENT_GROUP")
    ap.add_argument("--course", default=os.getenv("COURSE_NAME", ""), help="default: $COURSE_NAME")
    a = ap.parse_args()
    missing = [n for n, v in (("student (STUDENT_NAME)", a.student), ("group (STUDENT_GROUP)", a.group),
                              ("course (COURSE_NAME)", a.course)) if not v]
    if missing:
        sys.exit("Missing header data: " + ", ".join(missing) + " — pass --student/--group/--course or set the env vars.")
    sol = json.load(open(a.solution, encoding="utf8"))
    bad = 0
    if a.run:
        bad = run_all(sol)
        json.dump(sol, open(a.solution, "w", encoding="utf8"), ensure_ascii=False, indent=1)
    number = f"{a.week} – {sol['title']}" if sol.get("title") else a.week
    out = a.out or os.path.join(os.path.expanduser("~/Desktop"), f"Assignment_{a.week}.docx")
    build_docx(sol, out, a.student, a.group, a.course, number, a.date)
    print("Saved:", out)
    if bad:
        sys.exit(f"{bad} task(s) failed to run")
    if a.pdf:
        build_pdf(sol, a.pdf, a.student, a.group, a.course, number, a.date)
        print("Saved:", a.pdf)


if __name__ == "__main__":
    main()
