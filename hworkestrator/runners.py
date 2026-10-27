"""Run the skills: cppwork (deterministic); latexwork, jupiterwork and secwork (Claude solves, we build/run)."""
import html
import json
import re
import subprocess
import sys
from pathlib import Path

from . import pricing
from .config import settings

SKILLS_DIR = Path(__file__).parent / "skills"

LATEX_SYSTEM = r"""You solve a university assignment for a GTU student and
write the solution as the BODY of a LaTeX document (no preamble, no \begin{document}).
Rules:
- Georgian prose (natural Georgian), formulas in LaTeX math. XeLaTeX with FreeSerif is used.
- For every task in the assignment write:
    \problem{N}{<the task statement, in the assignment's own wording>}
    \answer
    <full solution: key steps, computations, final result>
  If the assignment contains a single task, use one \problem and organise the answer into parts.
- Available packages: amsmath, amssymb, mathtools, cancel, enumitem, booktabs, array, tikz.
  Georgian words inside math: \text{...}. Use pmatrix/align*/cases/tabular as needed.
- Be mathematically careful: double-check every arithmetic result. If the task is ambiguous,
  choose a concrete reasonable option, say so in one sentence and use it consistently
  (the student's name is {student} unless told otherwise).
- Output ONLY the LaTeX body, no markdown fences, no commentary."""


def _run(cmd, **kw):
    r = subprocess.run(cmd, capture_output=True, text=True, **kw)
    if r.returncode:
        lines = [l for l in (r.stdout + r.stderr).strip().splitlines() if l.strip()]
        raise RuntimeError(lines[-1] if lines else f"{Path(cmd[1]).name} failed")
    return r.stdout


def result_name(week_label):
    return f"დავალება {week_label} - {settings.student_name}"


# ---------- cppwork ----------
def header_args(course):
    """Student header data for the skills' build scripts (comes from .env, never from code)."""
    return ["--student", settings.student_name, "--group", settings.student_group, "--course", course]


def run_cppwork(pptx, week_label, outdir, course, stdin=None):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    out = outdir / f"{result_name(week_label)}.docx"
    cmd = [sys.executable, str(SKILLS_DIR / "cppwork/scripts/build.py"), str(pptx),
           "--week", week_label, "--out", str(out), *header_args(course)]
    for item in stdin or []:
        cmd += ["--stdin", item]
    log = _run(cmd)
    return out, log


# ---------- text extraction ----------
def extract_text(path):
    path = Path(path)
    suf = path.suffix.lower()
    if suf == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if suf == ".pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    if suf == ".pptx":
        import zipfile
        z = zipfile.ZipFile(path)
        names = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                       key=lambda n: int(re.findall(r"\d+", n)[0]))
        out = []
        for n in names:
            x = z.read(n).decode("utf8")
            out += [html.unescape("".join(re.findall(r"<a:t>([^<]*)</a:t>", p)))
                    for p in re.findall(r"<a:p>.*?</a:p>|<a:p [^>]*>.*?</a:p>", x, flags=re.S)]
        return "\n".join(t for t in out if t.strip())
    if suf in (".txt", ".md", ".tex"):
        return path.read_text(encoding="utf8", errors="ignore")
    raise ValueError(f"unsupported assignment file type: {path.name}")


PLACEHOLDERS = {"შეიტანეთ დავალების ტექსტი"}   # Moodle filler text on empty assignments


def html_to_text(s):
    from bs4 import BeautifulSoup
    text = BeautifulSoup(html.unescape(s or ""), "html.parser").get_text("\n").strip()
    return "" if text.rstrip(". ") in PLACEHOLDERS else text


# ---------- latexwork ----------
def _claude(system, user, max_tokens=12000, model=None):
    """One Claude call (streamed, so long outputs are allowed). Refuses a truncated answer."""
    import anthropic
    model = model or settings.anthropic_model
    with anthropic.Anthropic().messages.stream(
            model=model, max_tokens=max_tokens, system=system,
            messages=[{"role": "user", "content": user}]) as stream:
        msg = stream.get_final_message()
    pricing.record(model, msg.usage)                       # for the cost report
    if msg.stop_reason == "max_tokens":
        raise RuntimeError("Claude-ის პასუხი ძალიან გრძელი გამოვიდა და შეწყდა (max_tokens)")
    return "".join(b.text for b in msg.content if b.type == "text").strip()


def _strip_fences(text):
    return re.sub(r"^```[a-zA-Z]*\s*|\s*```$", "", text.strip())


def solve_to_latex(task_text, week_label, model=None):  # noqa: week_label kept for future prompts
    body = _strip_fences(_claude(LATEX_SYSTEM.replace("{student}", settings.student_name or "the student"),
                                 f"Assignment text:\n\n{task_text}", model=model))
    if "\\problem" not in body:
        raise RuntimeError("model output has no \\problem — unexpected format")
    return body


def run_latexwork(task_text, week_label, outdir, course, model=None):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    body = outdir / "body.tex"
    body.write_text(solve_to_latex(task_text, week_label, model), encoding="utf8")
    folder = outdir / "latex_result"
    name = f"Assignment_{week_label}"
    log = _run([sys.executable, str(SKILLS_DIR / "latexwork/scripts/build.py"), str(body),
                "--week", week_label, "--name", name, "--outdir", str(folder), *header_args(course)])
    pdf = folder / f"{name}.pdf"
    final = outdir / f"{result_name(week_label)}.pdf"
    final.write_bytes(pdf.read_bytes())
    zipped = Path(__import__("shutil").make_archive(str(outdir / f"{name}_tex"), "zip", folder))
    return final, zipped, log


# ---------- jupiterwork ----------
JUPYTER_SYSTEM = r"""You solve a data-science / machine-learning university assignment and write the
solution as a Jupyter notebook in jupytext "percent" format (NO header — it is added automatically).
Format:
    # %% [markdown]
    # ## დავალება 1. <title>
    # <task statement in the assignment's wording>

    # %%
    <python code>

    # %% [markdown]
    # ### პასუხი
    # TODO
Rules:
- Every markdown line starts with "# ". Georgian prose for headings/explanations; code and comments in English.
- Per task: title+statement, code cell(s) that really run, then the "### პასუხი" cell containing just the
  line "# TODO" — the Georgian explanations are written in a later step, from the real outputs, so do NOT
  write them now (but make the code print everything the questions of the task need).
- Datasets must work OFFLINE: use sklearn built-ins (Iris, Wine, Breast Cancer, Diabetes). Do not use Titanic /
  California Housing / any CSV that needs a download or does not exist. Keep one dataset per task family.
- random_state=42; print the metrics asked (accuracy, precision, recall, F1, MAE, MSE, RMSE, R²).
- Plots: matplotlib/seaborn, English labels, always plt.show().
- The notebook is executed top to bottom: imports in the first code cell, no input(), no files.
- Output ONLY the percent-format source: no markdown fences, no commentary."""


def expected_tasks(task_text):
    """Task numbers named in the assignment, e.g. 'დავალება 7.' -> 7."""
    return sorted({int(n) for n in re.findall(pricing.TASK_WORDS + r"\s*(\d+)\s*[.:)]", task_text, flags=re.I)})


def missing_tasks(cells_text, task_text):
    have = {int(n) for n in re.findall(r"^#\s*#+\s*" + pricing.TASK_WORDS + r"\s*(\d+)", cells_text,
                                       flags=re.M | re.I)}
    return [n for n in expected_tasks(task_text) if n not in have]


EXPLAIN_SYSTEM = r"""You write the explanations ("პასუხი") of a data-science notebook from its REAL outputs.
Return ONLY a JSON object {"1": "<markdown>", "2": "<markdown>", ...}: one entry per task number, the
markdown being the Georgian answer that fully answers every question/bullet of that task and cites the
actual numbers from the outputs. No code fences, no commentary, valid JSON (escape newlines as \\n)."""


def splice_answers(cells_text, answers):
    """Replace the markdown after '### პასუხი' of task N with answers[N] (percent-format source)."""
    lines = cells_text.split("\n")
    for n, md in answers.items():
        head = next((i for i, l in enumerate(lines) if re.match(rf"^#\s*#+\s*{pricing.TASK_WORDS}\s*{n}\b", l, re.I)), None)
        if head is None:
            raise ValueError(f"task {n} not found")
        ans = next((i for i in range(head, len(lines)) if re.match(r"^#\s*#+\s*პასუხი", lines[i])), None)
        if ans is None:
            raise ValueError(f"answer heading of task {n} not found")
        end = next((i for i in range(ans + 1, len(lines)) if lines[i].startswith("# %%")), len(lines))
        body = [("# " + l) if l.strip() else "#" for l in str(md).strip().split("\n")]
        lines[ans + 1:end] = body + [""]
    return "\n".join(lines)


def _jupyter_build(cells_text, folder, name, week_label, course):
    cells = Path(folder) / "cells.py"
    Path(folder).mkdir(parents=True, exist_ok=True)
    cells.write_text(cells_text, encoding="utf8")
    cmd = [sys.executable, str(SKILLS_DIR / "jupiterwork/scripts/build.py"), str(cells),
           "--week", f"{week_label} კვირა", "--course", course, "--name", name, "--outdir", str(folder),
           "--student", settings.student_name, "--group", settings.student_group]
    r = subprocess.run(cmd, capture_output=True, text=True)
    return r.returncode == 0, r.stdout + r.stderr


def run_jupiterwork(task_text, week_label, outdir, course, model=None):
    """Claude writes the notebook source, we execute it (repairing errors), then Claude writes the
    explanations from the REAL outputs and we execute again."""
    outdir = Path(outdir)
    folder, name = outdir / "nb_result", f"Assignment_{week_label}"
    cells = _strip_fences(_claude(JUPYTER_SYSTEM, f"Assignment text:\n\n{task_text}", 32000, model))
    for _ in range(2):                                    # make sure no task was skipped
        gone = missing_tasks(cells, task_text)
        if not gone:
            break
        cells = _strip_fences(_claude(
            JUPYTER_SYSTEM, f"Assignment text:\n\n{task_text}\n\nYour previous source:\n{cells}\n\n"
            f"It is missing tasks {gone}. Return the FULL source with ALL tasks.", 32000, model))
    ok, log = _jupyter_build(cells, folder, name, week_label, course)
    for _ in range(2):                                    # repair loop
        if ok:
            break
        cells = _strip_fences(_claude(
            JUPYTER_SYSTEM, f"Assignment text:\n\n{task_text}\n\nYour previous notebook source:\n{cells}\n\n"
            f"Executing it failed with:\n{log[-2500:]}\n\nReturn the full corrected source.", 32000, model))
        ok, log = _jupyter_build(cells, folder, name, week_label, course)
    if not ok:
        raise RuntimeError("notebook execution failed: " + (log.strip().splitlines() or ["?"])[-1])
    outputs = log.split("=== cell outputs", 1)[-1][:9000]
    ok2, log2, revised = False, log, cells
    for _ in range(2):                                    # pass 2: the explanations (one retry)
        try:
            raw = _strip_fences(_claude(
                EXPLAIN_SYSTEM, f"Assignment text:\n\n{task_text}\n\nNotebook source:\n{cells}\n\n"
                f"REAL outputs of running it:\n{outputs}", 16000, model))
            answers = {int(k): v for k, v in json.loads(raw).items()}
            if set(answers) < set(expected_tasks(task_text)):
                raise ValueError("answers missing")
            revised = splice_answers(cells, answers)
            ok2, log2 = _jupyter_build(revised, folder, name, week_label, course)
            if ok2 and "# TODO" not in revised:
                break
        except Exception:
            ok2 = False
    if not ok2 or "# TODO" in revised:                    # never ship a notebook with "TODO" answers
        raise RuntimeError("ახსნების (პასუხების) გენერაცია ვერ მოხერხდა — სცადე თავიდან ან სხვა მოდელით")
    final = outdir / f"{result_name(week_label)}.ipynb"
    final.write_bytes((folder / f"{name}.ipynb").read_bytes())
    return final, [], log2 if ok2 else log


# ---------- secwork (information security) ----------
SEC_SYSTEM = r"""You solve an information-security university assignment (ciphers, encodings, hashing,
number systems, probability of breaking a key, security theory...) for a GTU student.
Return ONLY a JSON object:
{"title": "<short Georgian title of the whole assignment, e.g. სახელისა და გვარის დაშიფვრა და ანალიზი>",
 "tasks": [{"n": 1, "title": "<short Georgian title>",
            "condition": ["<what this task does, in the assignment's wording>", "<keys / method / formulas>"],
            "code": "<complete Python 3 program, or \"\" for a theory-only task>",
            "caption": "<Georgian caption of the output, e.g. სახელისა და გვარის დაშიფვრის შედეგი>",
            "answer": ["<Georgian paragraph>", ...]}],
 "conclusions": []}
Rules:
- Split the assignment into tasks: numbered items as they are; an unnumbered list of steps ("encrypt...",
  "convert...", "analyse...") -> one task per step, in order. Never skip a step. Natural Georgian prose.
  Formulas as plain text with Unicode (C = (P + k) mod 26, P ⊕ K, ≡, ⁻¹, ≈), no LaTeX.
- "condition" states the concrete choices: the keys used (fixed, e.g. Caesar k = 7, Vigenère CIPHER, Vernam key
  of the text's length), how non-letters are handled, the formulas.
- A task that asks to encrypt/decrypt/compute/convert/compare/analyse gets a "code": standard library only,
  no input(), no files, no network, deterministic (fixed keys, or random.seed(42) and print the key).
  Every task's code is run on its own, so it is self-contained: a task that uses an earlier result recomputes
  it (or uses the same fixed values) — but keep it to what THAT task needs and print only its own results.
  Print every value the task asks for with aligned Georgian labels, e.g. print(f"{'ღია ტექსტი':<12}: {p}").
  Comments in Georgian. Byte values: hex as "0A 1F ...", binary 8 digits, octal 3 digits.
  For such a task leave "answer" as [] — the analysis is written later from the REAL output.
- A theory task gets "code": "" and a full Georgian "answer".
- Personal data: the student is {student}; when a cipher needs Latin letters, transliterate the name
  (e.g. ანდრია გვარამია -> Andria Gvaramia) and say so in the condition.
- Valid JSON only (escape newlines in code as \n), no markdown fences, no commentary."""

SEC_EXPLAIN_SYSTEM = r"""You write the analysis of an information-security assignment from the REAL outputs
of its programs. Return ONLY JSON:
{"answers": {"<task n>": ["<Georgian paragraph>", ...]}, "conclusions": ["<Georgian paragraph>", ...]}
"answers": ONLY for a task whose question is not answered by its output alone (e.g. "which is stronger and
why"); omit the others. "conclusions": numbered paragraphs "1. <method/aspect>. ..." — one per method or aspect
(what the output shows, key space, weaknesses: frequency analysis, brute force, key reuse...), each citing the
actual numbers from the outputs; then a last paragraph starting "საერთო დასკვნა:". Valid JSON, no fences."""


def _json(raw):
    raw = _strip_fences(raw)
    return json.loads(raw[raw.find("{"):raw.rfind("}") + 1])


def _sec_build(sol, folder, week_label, course, run, pdf=False):
    path = Path(folder) / "solution.json"
    Path(folder).mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(sol, ensure_ascii=False, indent=1), encoding="utf8")
    cmd = [sys.executable, str(SKILLS_DIR / "secwork/scripts/build.py"), str(path), "--week", week_label,
           "--out", str(Path(folder) / "out.docx"), *header_args(course)]
    cmd += ["--run"] if run else []
    cmd += ["--pdf", str(Path(folder) / "out.pdf")] if pdf else []
    r = subprocess.run(cmd, capture_output=True, text=True)
    return r.returncode == 0, r.stdout + r.stderr, json.loads(path.read_text(encoding="utf8"))


def run_secwork(task_text, week_label, outdir, course, model=None):
    """Claude writes tasks + Python code, we run the code (repairing errors), then Claude explains the
    REAL outputs and writes the conclusions."""
    outdir = Path(outdir)
    folder = outdir / "sec_result"
    system = SEC_SYSTEM.replace("{student}", settings.student_name or "the student")
    user = f"Assignment text:\n\n{task_text}"
    sol = _json(_claude(system, user, 24000, model))
    gone = [n for n in expected_tasks(task_text) if n not in {int(t["n"]) for t in sol["tasks"]}]
    if gone:                                              # make sure no task was skipped
        sol = _json(_claude(system, f"{user}\n\nYour previous answer:\n{json.dumps(sol, ensure_ascii=False)}\n\n"
                            f"It is missing tasks {gone}. Return the FULL JSON with ALL tasks.", 24000, model))
    ok, log, sol = _sec_build(sol, folder, week_label, course, run=True)
    for _ in range(2):                                    # repair loop
        if ok:
            break
        sol = _json(_claude(system, f"{user}\n\nYour previous answer:\n{json.dumps(sol, ensure_ascii=False)}\n\n"
                            f"Running the code failed with:\n{log[-2500:]}\n\nReturn the full corrected JSON.",
                            24000, model))
        ok, log, sol = _sec_build(sol, folder, week_label, course, run=True)
    if not ok:
        raise RuntimeError("კოდის გაშვება ვერ მოხერხდა: " + (log.strip().splitlines() or ["?"])[-1])
    coded = [t for t in sol["tasks"] if t.get("code", "").strip()]
    if coded:                                             # pass 2: explanations from the real outputs
        shown = "\n\n".join(f"Task {t['n']} code:\n{t['code']}\n\nREAL output:\n{t.get('output', '')[:3000]}"
                            for t in coded)
        ex = _json(_claude(SEC_EXPLAIN_SYSTEM, f"Assignment text:\n\n{task_text}\n\n{shown}", 12000, model))
        answers = {int(k): v for k, v in (ex.get("answers") or {}).items()}
        for t in coded:
            t["answer"] = answers.get(int(t["n"])) or t.get("answer") or []
        sol["conclusions"] = ex.get("conclusions") or sol.get("conclusions") or []
    ok, log, sol = _sec_build(sol, folder, week_label, course, run=False, pdf=True)   # keeps the stored outputs
    if not ok:
        raise RuntimeError("PDF/docx build failed: " + (log.strip().splitlines() or ["?"])[-1])
    pdf, docx = outdir / f"{result_name(week_label)}.pdf", outdir / f"{result_name(week_label)}.docx"
    pdf.write_bytes((folder / "out.pdf").read_bytes())
    docx.write_bytes((folder / "out.docx").read_bytes())
    return pdf, [docx], log          # the PDF is uploaded; the .docx is sent too, for manual edits


# ---------- price estimation helper (uses the cheapest model) ----------
SIZE_SYSTEM = r"""You help estimate the workload of a university assignment before it is solved.
Return ONLY JSON: {"tasks": [{"n": 1, "size": "S"}, ...]} — one entry per numbered task/problem.
size: "S" = trivial or one short answer, "M" = typical (some code or calculation + a paragraph),
"L" = heavy (model training, several plots/steps, long multi-part questions). No commentary."""


def judge_task_sizes(task_text):
    """Ask Haiku 5.5 (a fraction of a cent) how big each task is -> ["S"/"M"/"L", ...]. [] on failure."""
    try:
        haiku = pricing.MODELS["haiku"]["id"]
        raw = _strip_fences(_claude(SIZE_SYSTEM, task_text[:20000], 1000, haiku))
        sizes = [t.get("size", "M") for t in json.loads(raw)["tasks"]]
        return [z if z in ("S", "M", "L") else "M" for z in sizes]
    except Exception:
        return []
