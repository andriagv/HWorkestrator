"""Reading assignments: text out of .docx / .pdf / .pptx / Moodle HTML, and finding the numbered tasks in it."""
import html
import re
from pathlib import Path

# words that introduce a numbered task: "დავალება 3.", "Task 2:", "Problem 4)"
TASK_WORDS = r"(?:დავალება|ამოცანა|პრობლემა|Task|Problem|Exercise|Question)"
PLACEHOLDERS = {"შეიტანეთ დავალების ტექსტი"}   # Moodle filler text on empty assignments


def extract_text(path):
    path = Path(path)
    suf = path.suffix.lower()
    if suf == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if suf == ".pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    if suf == ".pptx":
        import zipfile
        z = zipfile.ZipFile(path)
        names = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                       key=lambda n: int(re.findall(r"\d+", n)[0]))
        out = []
        for n in names:
            x = z.read(n).decode("utf8")
            out += [html.unescape("".join(re.findall(r"<a:t>([^<]*)</a:t>", p)))
                    for p in re.findall(r"<a:p>.*?</a:p>|<a:p [^>]*>.*?</a:p>", x, flags=re.S)]
        return "\n".join(t for t in out if t.strip())
    if suf in (".txt", ".md", ".tex"):
        return path.read_text(encoding="utf8", errors="ignore")
    raise ValueError(f"unsupported assignment file type: {path.name}")


def html_to_text(s):
    from bs4 import BeautifulSoup
    text = BeautifulSoup(html.unescape(s or ""), "html.parser").get_text("\n").strip()
    return "" if text.rstrip(". ") in PLACEHOLDERS else text


def expected_tasks(task_text):
    """Task numbers named in the assignment, e.g. 'დავალება 7.' -> 7."""
    return sorted({int(n) for n in re.findall(TASK_WORDS + r"\s*(\d+)\s*[.:)]", task_text, flags=re.I)})


def n_tasks(task_text):
    return max(len(expected_tasks(task_text)), 1)
