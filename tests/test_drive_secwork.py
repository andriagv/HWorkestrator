"""Info-security flow: assignment file in the Drive folder -> secwork (Claude mocked, code really runs, PDF via
XeLaTeX) -> upload into a folder the user picks."""
import json
import shutil
from pathlib import Path

import pytest

from hworkestrator import gdrive, pipeline, runners

FILES = {
    "f3": {"id": "f3", "name": "დავალება 3.pdf", "mimeType": "application/pdf"},
    "f4": {"id": "f4", "name": "დავალება 4", "mimeType": gdrive.DOCX},          # uploaded without extension
    "form": {"id": "form", "name": "გამოკითხვა", "mimeType": "application/vnd.google-apps.form"},
}
FOLDERS = {
    "root": [FILES["f3"], FILES["f4"], FILES["form"],
             {"id": "d4", "name": "4 დავ. მიმდინარე", "mimeType": gdrive.GFOLDER},
             {"id": "d3", "name": "3 დავალება", "mimeType": gdrive.GFOLDER}],
    "d3": [{"id": "g1", "name": "000000-0", "mimeType": gdrive.GFOLDER},
           {"id": "g2", "name": "111111-1", "mimeType": gdrive.GFOLDER}],
    "g1": [{"id": "x", "name": "old.pdf", "mimeType": "application/pdf"}],
}
NAMES = {"d3": "3 დავალება", "d4": "4 დავ. მიმდინარე", "g1": "000000-0", **{k: v["name"] for k, v in FILES.items()}}
TASK = "ამოცანა 1. დაშიფრეთ HELLO ცეზარის მეთოდით (k = 3).\nამოცანა 2. რა არის CIA ტრიადა?"


class FakeDrive(gdrive.Drive):
    def __init__(self):
        super().__init__("cid", "s", "r")
        self.uploaded = None

    def children(self, parent_id, order=None):
        return FOLDERS[parent_id]

    def get(self, file_id):
        return FILES.get(file_id) or {"id": file_id, "name": NAMES[file_id], "mimeType": gdrive.GFOLDER}

    def download(self, f, dest_dir):
        assert gdrive.downloadable(f)
        Path(dest_dir).mkdir(parents=True, exist_ok=True)
        path = Path(dest_dir) / (Path(f["name"]).stem + ".txt")
        path.write_text(TASK, encoding="utf8")
        return path

    def upload(self, folder_id, path, name):
        self.uploaded = (folder_id, name)


SOLUTION = {"title": "ცეზარის შიფრი", "tasks": [
    {"n": 1, "title": "ცეზარი", "condition": ["C = (P + k) mod 26"], "caption": "შედეგი", "answer": [],
     "code": "print('შიფრი :', ''.join(chr((ord(c) - 65 + 3) % 26 + 65) for c in 'HELLO'))"},
    {"n": 2, "title": "CIA", "condition": ["რა არის CIA?"], "code": "", "answer": ["კონფიდენციალურობა..."]}],
    "conclusions": []}


@pytest.fixture
def fake(monkeypatch, tmp_path):
    drv = FakeDrive()
    monkeypatch.setattr(pipeline, "_drive", lambda: drv)
    monkeypatch.setattr(pipeline, "new_workdir", lambda: tmp_path)
    return drv


def test_drive_lists_assignment_files_and_folders(fake):
    assert pipeline.list_items("infosec") == [("დავალება 3.pdf", "f3"), ("დავალება 4", "f4")]   # no form, no folders
    ov = pipeline.overview("infosec", "f3")
    assert ov.label == "3" and [f.name for f in ov.files] == ["დავალება 3.txt"]
    assert pipeline.overview("infosec", "f4").label == "4"
    root = pipeline.drive_folder()
    assert root["is_root"] and [n for n, _ in root["folders"]] == ["3 დავალება", "4 დავ. მიმდინარე"]
    d3 = pipeline.drive_folder("d3")
    assert d3["name"] == "3 დავალება" and [n for n, _ in d3["folders"]] == ["000000-0", "111111-1"]


@pytest.mark.skipif(not shutil.which("xelatex"), reason="xelatex not installed")
def test_secwork_on_drive_end_to_end(fake, monkeypatch):
    def claude(system, user, max_tokens=0, model=None):
        if system is runners.SEC_EXPLAIN_SYSTEM:
            assert "KHOOR" in user                       # the analysis gets the REAL output
            return json.dumps({"answers": {}, "conclusions": ["1. ცეზარი სუსტია.", "საერთო დასკვნა: ..."]})
        return "```json\n" + json.dumps(SOLUTION, ensure_ascii=False) + "\n```"
    monkeypatch.setattr(runners, "_claude", claude)

    ov = pipeline.overview("infosec", "f3")
    job = pipeline.prepare(ov, "secwork")
    assert job.upload_file.name == "დავალება 3 - Test Student.pdf"
    assert job.upload_file.read_bytes()[:4] == b"%PDF"
    assert [f.name for f in job.extra_files] == ["დავალება 3 - Test Student.docx"]
    import docx
    d = docx.Document(str(job.extra_files[0]))
    texts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for c in t._cells]
    assert any("შიფრი : KHOOR" in t for t in texts)                    # real program output
    assert "დასკვნები" in texts and "კონფიდენციალურობა..." in texts
    assert any("3 – ცეზარის შიფრი" in t for t in texts)                # "დავალების ნომერი" in the header
    with pytest.raises(pipeline.PipelineError):                        # no folder chosen yet
        pipeline.submit(job)
    job.target = {"kind": "drive", "folder_id": "g1", "label": "000000-0"}
    assert "000000-0" in pipeline.submit(job)
    assert fake.uploaded == ("g1", "დავალება 3 - Test Student.pdf")


def test_secwork_repairs_failing_code(monkeypatch, tmp_path):
    broken = json.loads(json.dumps(SOLUTION))
    broken["tasks"][0]["code"] = "print(1/0)"
    answers = iter([broken, SOLUTION])
    calls = []

    def claude(system, user, max_tokens=0, model=None):
        calls.append(user)
        if system is runners.SEC_EXPLAIN_SYSTEM:
            return '{"answers": {}, "conclusions": []}'
        return json.dumps(next(answers), ensure_ascii=False)
    monkeypatch.setattr(runners, "_claude", claude)
    monkeypatch.setattr(runners, "_sec_build", _no_pdf(runners._sec_build))
    out, extra, _ = runners.run_secwork(TASK, "2", tmp_path, "ინფ. უსაფრთხოება")
    assert "ZeroDivisionError" in calls[1]


def _no_pdf(build):
    """Skip XeLaTeX (only the code-repair loop is tested); fake the PDF from the docx."""
    def wrapped(sol, folder, week_label, course, run, pdf=False):
        ok, log, sol = build(sol, folder, week_label, course, run)
        if pdf:
            (Path(folder) / "out.pdf").write_bytes(b"%PDF-fake")
        return ok, log, sol
    return wrapped


def test_downloaded_file_without_extension_gets_one(tmp_path):
    class S:
        def get(self, url, **kw):
            return type("R", (), {"status_code": 200, "content": b"x", "json": lambda self: {}})()
    drv = gdrive.Drive("c", "s", "r", S())
    drv._access = "a"
    assert drv.download(FILES["f4"], tmp_path).name == "დავალება 4.docx"
    assert drv.download(FILES["f3"], tmp_path).name == "დავალება 3.pdf"


class Resp:
    status_code = 200

    def __init__(self, data=None):
        self.data = data or {}

    def json(self):
        return self.data


class Session:
    def __init__(self, existing):
        self.existing, self.sent = existing, []

    def post(self, url, **kw):
        self.sent.append(("POST", url, kw.get("params")))
        return Resp({"access_token": "a"} if "oauth2" in url else {"id": "new"})

    def patch(self, url, **kw):
        self.sent.append(("PATCH", url, kw.get("params")))
        return Resp({"id": url.rsplit("/", 1)[1]})

    def get(self, url, **kw):
        return Resp({"files": self.existing})


def test_upload_replaces_a_file_with_the_same_name(tmp_path):
    f = tmp_path / "a.pdf"
    f.write_bytes(b"x")
    s = Session([{"id": "f1", "name": "დავალება 3 - X.pdf", "mimeType": "application/pdf"}])
    assert gdrive.Drive("c", "s", "r", s).upload("dir", f, "დავალება 3 - X.pdf") == "f1"
    assert s.sent[-1][0] == "PATCH"
    s = Session([])
    assert gdrive.Drive("c", "s", "r", s).upload("dir", f, "დავალება 3 - X.pdf") == "new"
    assert s.sent[-1][0] == "POST"
