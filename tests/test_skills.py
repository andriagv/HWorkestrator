"""Skills, through the pipeline: cppwork on a fake Moodle, secwork on a fake Drive (Claude mocked, the code and
the build really run), the jupiterwork helpers, and the generic checks in pipeline.prepare."""
import json
import shutil
from pathlib import Path

import pytest
from fakes import TASK, FakeMoodle

from hworkestrator import llm, pipeline, skills
from hworkestrator.documents import expected_tasks
from hworkestrator.sources import moodle as moodle_src

LECTURE = Path.home() / "Downloads" / "Lecture IV_struct.pptx"
secwork = skills.load("secwork")
jupiterwork = skills.load("jupiterwork")


# ---------------------------------------------------------------- pipeline.prepare
@pytest.mark.skipif(not LECTURE.exists(), reason="sample lecture not available")
def test_cppwork_moodle_end_to_end(monkeypatch, tmp_path):
    fake = FakeMoodle(lecture=LECTURE)
    monkeypatch.setattr(moodle_src, "client", lambda: fake)
    monkeypatch.setattr(pipeline, "new_workdir", lambda: tmp_path)
    ov = pipeline.overview("dsa", "4")
    job = pipeline.prepare(ov, "cppwork")
    assert job.upload_file.suffix == ".docx" and job.upload_file.exists()
    assert job.target == {"kind": "moodle", "assignment_id": 999, "label": "დავალება 4", "existing": ["old.docx"]}
    monkeypatch.setattr(pipeline.settings.__class__, "moodle_finalize", False, raising=False)
    pipeline.submit(job)
    assert fake.submitted == (999, "დავალება IV - Test Student.docx")


def test_a_lecture_skill_needs_a_lecture_file_among_chosen_files(tmp_path):
    ov = pipeline.Overview("dsa", 4, "t", "x", files=[tmp_path / "notes.txt"], intro="y" * 30,
                           targets=[{"kind": "moodle", "assignment_id": 1}])
    with pytest.raises(pipeline.PipelineError):                   # .txt is not a lecture
        pipeline.prepare(ov, "cppwork", chosen=[tmp_path / "notes.txt"])


def test_unknown_skill_is_rejected(tmp_path):
    ov = pipeline.Overview("dsa", 4, "t", "x", files=[], intro="y" * 30)
    with pytest.raises(pipeline.PipelineError):
        pipeline.prepare(ov, "nope")


def test_every_skill_follows_the_contract():
    from hworkestrator.registry import SKILLS
    for name in SKILLS:
        r = skills.load(name)
        assert r.NEEDS in ("text", "lecture") and callable(r.run)
        if r.USES_LLM:
            calls = r.estimate(1000, ["M", "L"])
            assert calls and all(i > 0 and o > 0 for i, o in calls)


# ---------------------------------------------------------------- secwork
SOLUTION = {"title": "ცეზარის შიფრი", "tasks": [
    {"n": 1, "title": "ცეზარი", "condition": ["C = (P + k) mod 26"], "caption": "შედეგი", "answer": [],
     "code": "print('შიფრი :', ''.join(chr((ord(c) - 65 + 3) % 26 + 65) for c in 'HELLO'))"},
    {"n": 2, "title": "CIA", "condition": ["რა არის CIA?"], "code": "", "answer": ["კონფიდენციალურობა..."]}],
    "conclusions": []}


@pytest.mark.skipif(not shutil.which("xelatex"), reason="xelatex not installed")
def test_secwork_on_drive_end_to_end(drive, monkeypatch):
    def claude(system, user, max_tokens=0, model=None):
        if system is secwork.EXPLAIN_SYSTEM:
            assert "KHOOR" in user                       # the analysis gets the REAL output
            return json.dumps({"answers": {}, "conclusions": ["1. ცეზარი სუსტია.", "საერთო დასკვნა: ..."]})
        return "```json\n" + json.dumps(SOLUTION, ensure_ascii=False) + "\n```"
    monkeypatch.setattr(llm, "claude", claude)

    job = pipeline.prepare(pipeline.overview("infosec", "f3"), "secwork")
    assert job.upload_file.name == "დავალება 3 - Test Student.pdf"
    assert job.upload_file.read_bytes()[:4] == b"%PDF"
    assert [f.name for f in job.extra_files] == ["დავალება 3 - Test Student.docx"]
    import docx
    d = docx.Document(str(job.extra_files[0]))
    texts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for c in t._cells]
    assert any("შიფრი : KHOOR" in t for t in texts)                    # real program output
    assert "დასკვნები" in texts and "კონფიდენციალურობა..." in texts
    assert any("3 – ცეზარის შიფრი" in t for t in texts)                # "დავალების ნომერი" in the header
    with pytest.raises(pipeline.PipelineError):                        # no folder chosen yet
        pipeline.submit(job)
    job.target = {"kind": "drive", "folder_id": "g1", "label": "000000-0"}
    assert "000000-0" in pipeline.submit(job)
    assert drive.uploaded == ("g1", "დავალება 3 - Test Student.pdf")


def test_secwork_repairs_failing_code(monkeypatch, tmp_path):
    broken = json.loads(json.dumps(SOLUTION))
    broken["tasks"][0]["code"] = "print(1/0)"
    answers, calls = iter([broken, SOLUTION]), []

    def claude(system, user, max_tokens=0, model=None):
        calls.append(user)
        if system is secwork.EXPLAIN_SYSTEM:
            return '{"answers": {}, "conclusions": []}'
        return json.dumps(next(answers), ensure_ascii=False)
    monkeypatch.setattr(llm, "claude", claude)

    real_build = secwork.build

    def no_pdf(sol, folder, label, course, run_code, pdf=False):   # only the repair loop is tested here
        ok, log, sol = real_build(sol, folder, label, course, run_code)
        if pdf:
            (Path(folder) / "out.pdf").write_bytes(b"%PDF-fake")
        return ok, log, sol
    monkeypatch.setattr(secwork, "build", no_pdf)
    secwork.run(TASK, "2", tmp_path, "ინფ. უსაფრთხოება")
    assert "ZeroDivisionError" in calls[1]


# ---------------------------------------------------------------- jupiterwork helpers
def test_task_completeness_check():
    text = "დავალება 1. ა\nდავალება 2. ბ\nდავალება 3: გ"
    assert expected_tasks(text) == [1, 2, 3]
    cells = "# %% [markdown]\n# ## დავალება 1. ა\n# %% [markdown]\n# ## დავალება 3. გ\n"
    assert jupiterwork.missing_tasks(cells, text) == [2]


def test_english_and_georgian_task_labels_are_counted():
    text = "Task 1: a\nTask 2: b\nამოცანა 3. გ\nProblem 4) d"
    assert expected_tasks(text) == [1, 2, 3, 4]
    cells = "# %% [markdown]\n# ## Task 1: a\n# %% [markdown]\n# ## Task 2: b\n"
    assert jupiterwork.missing_tasks(cells, text) == [3, 4]


def test_splice_answers_replaces_only_the_answer_blocks():
    src = ("# %% [markdown]\n# ## დავალება 1. ა\n# %%\nprint(1)\n# %% [markdown]\n# ### პასუხი\n# old1\n"
           "# %% [markdown]\n# ## დავალება 2. ბ\n# %%\nprint(2)\n# %% [markdown]\n# ### პასუხი\n# old2\n# more\n")
    out = jupiterwork.splice_answers(src, {1: "new **one**", 2: "line a\n\nline b"})
    assert "old1" not in out and "old2" not in out and "more" not in out
    assert "# new **one**" in out and "# line a\n#\n# line b" in out
    assert out.count("print(") == 2 and "## დავალება 2" in out
    with pytest.raises(ValueError):
        jupiterwork.splice_answers(src, {7: "x"})


# ---------------------------------------------------------------- cppwork: PDF / Word lectures
def _cppwork():
    import importlib.util, pathlib
    spec = importlib.util.spec_from_file_location(
        "cppwork_build", pathlib.Path(__file__).parent.parent / "hworkestrator/skills/cppwork/scripts/build.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def _mini_docx(path, paragraphs):
    import zipfile
    body = "".join(f"<w:p><w:r><w:t xml:space=\"preserve\">{p}</w:t></w:r></w:p>" for p in paragraphs)
    xml = ('<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f"<w:body>{body}</w:body></w:document>")
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", xml)


def test_cppwork_reads_word_lectures_marked_and_unmarked(tmp_path):
    b = _cppwork()
    marked = tmp_path / "m.docx"
    _mini_docx(marked, ["საილუსტრაციო მაგალითი_1: დაწერეთ კოდი,", "რომელიც ბეჭდავს ორ რიცხვს.",
                        "#include &lt;iostream&gt;", "using namespace std;", "int main()", "{",
                        "cout &lt;&lt; “hi” &lt;&lt; endl;", "return 0;", "}", "შედეგი: hi"])
    ex = b.extract(str(marked))
    assert len(ex) == 1 and "ორ რიცხვს" in ex[0]["cond"]                  # wrapped condition is joined
    code = "\n".join(b.tidy(ex[0]["code"]))
    assert 'cout << "hi" << endl;' in code and "“" not in code           # smart quotes fixed

    plain = tmp_path / "u.docx"
    _mini_docx(plain, ["სტეკი", "კოდი code", "#include &lt;iostream&gt;", "int main()", "{", "return 0;", "}",
                       "რიგი", "#include &lt;iostream&gt;", "int main()", "{", "return 0;", "}"])
    blocks = b.extract_unmarked(str(plain))
    assert [x["cond"].split(":")[0] for x in blocks] == ["სტეკი", "რიგი"]


def test_cppwork_reads_pdf_pages_and_drops_page_numbers(monkeypatch, tmp_path):
    b = _cppwork()
    import pypdf

    class FakePage:
        def __init__(self, t): self.t = t
        def extract_text(self): return self.t

    class FakeReader:
        def __init__(self, path):
            self.pages = [FakePage("საილუსტრაციო მაგალითი_3: ჯამი\n#include <iostream>\nint main()\n{\n3\nreturn 0;\n}\nშედეგი: 0\n12")]

    monkeypatch.setattr(pypdf, "PdfReader", FakeReader)
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF")
    ex = b.extract(str(pdf))
    assert len(ex) == 1 and ex[0]["n"] == 3
    assert ex[0]["code"] == ["#include <iostream>", "int main()", "{", "return 0;", "}"]   # lone "3" dropped


def test_cppwork_pdf_code_continues_over_a_page_break_and_keeps_its_layout(monkeypatch, tmp_path):
    b = _cppwork()
    import pypdf

    head = "ლექცია 6 | ალგორითმები"                       # running header on every page
    page1 = (head + "\n1\nSLL-ის სრული C++ კოდი\n#include <iostream>\nusing namespace std;\n\n"
             "struct Node {\n    int data;\n    Node(int v) : data(v) {}")
    page2 = (head + "\n2\n\n};\n\nint main() {\n    Node n(1);\n    cout << n.data << endl;\n"
             "    return 0;\n}\n\nდასკვნითი რეკომენდაციები")

    class FakePage:
        def __init__(self, t): self.t = t
        def extract_text(self): return self.t

    class FakeReader:
        def __init__(self, path): self.pages = [FakePage(page1), FakePage(page2)]

    monkeypatch.setattr(pypdf, "PdfReader", FakeReader)
    pdf = tmp_path / "lec.pdf"
    pdf.write_bytes(b"%PDF")
    ex = b.extract(str(pdf)) or b.extract_unmarked(str(pdf))
    assert len(ex) == 1 and ex[0]["cond"].startswith("SLL-ის სრული C++ კოდი")
    code = b.tidy(ex[0]["code"], verbatim=True)
    assert head not in "\n".join(code)                      # running header removed
    assert "    int data;" in code and code[-1] == "}"       # own indentation kept, block spans the page break
    assert code[code.index("};") - 1] != ""                  # no blank line right before a closing brace


def test_cppwork_copies_pdf_and_word_code_exactly(tmp_path):
    """Lecturers expect the code from a PDF/Word lecture 1:1: block comments, Georgian text after a
    statement and odd spacing must not be 'cleaned'."""
    b = _cppwork()
    doc = tmp_path / "x.docx"
    _mini_docx(doc, ["საილუსტრაციო მაგალითი_1: ტესტი", "#include &lt;iostream&gt;", "int main() {",
                     "  /* პირველი ხაზი", "     მეორე ხაზი */", "  int  x =  5;   // ხუთი",
                     "  cout &lt;&lt; x ;  ბეჭდავს", "  return 0;", "}"])
    ex = b.extract(str(doc))[0]
    code = b.tidy(ex["code"], verbatim=True)
    assert code == ["#include <iostream>", "int main() {", "  /* პირველი ხაზი", "     მეორე ხაზი */",
                    "  int  x =  5;   // ხუთი", "  cout << x ;  ბეჭდავს", "  return 0;", "}"]
